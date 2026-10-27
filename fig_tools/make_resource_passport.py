"""Build the downloadable Resource Passport worksheet for the Part 107 Studying lab.

Writes docs/labs/files/lab06_resource_passport.xlsx, one sheet per lab part. The rows must match
the Parts on docs/labs/6_part_107_studying.md; edit both together. Same plain style as
make_mission_log.py.

    python fig_tools/make_resource_passport.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "docs" / "labs" / "files" / "lab06_resource_passport.xlsx"

BLUE = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=BLUE)
LABEL_FILL = PatternFill("solid", fgColor="DDEBF7")
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# Part 1 - the ten finds: (resource, what to write)
FINDS = [
    ("Testing supplement FAA-CT-8080-2H (course PDF)", "The PDF page number of the sectional chart legend"),
    ("Supplement Figure 26 (PDF page 59)", "The name of the airport at area 2"),
    ("Supplement Figure 12 (PDF page 46)", "The station whose report is a SPECI with rain"),
    ("AC 107-2A (PDF)", "The chapter that covers the Declaration of Compliance"),
    ("eCFR, 14 CFR 107.51", "The maximum groundspeed, in knots"),
    ("eCFR, 14 CFR 107.9", "How many days you have to report a safety event"),
    ("PHAK chapter 15, Airspace (open the chapter PDF)", "The airspace class that begins at 18,000 ft MSL"),
    ("FAA LAANC page", "The name of one approved provider that offers public service"),
    ("PSI site, faa.psiexams.com", "The menu path to the free UAG sample test"),
    ("FAA UAG sample-question PDF", "How many questions it has, and whether it includes an answer key"),
]

# Part 2 - figure questions: (label, FAA sample question, figure cited)
FIGURE_QS = [
    ("Demo (TA)", "FAA sample Q9", "Figure 20, area 3"),
    ("1", "FAA sample Q24", "Figure 21"),
    ("2", "FAA sample Q26", "Figure 12"),
    ("3", "FAA sample Q31", "Figure 26, area 4"),
]

# Part 3 - the look-it-up race
RACE = [
    "Minimum flight visibility from the control station?",
    "Minimum distance below a cloud, and horizontally from it?",
    "How many hours after drinking alcohol before you may fly?",
    "Who submits a Declaration of Compliance, and for which categories of operation over people?",
    "A standard Remote ID aircraft stops broadcasting in flight. What must the person on the controls do?",
    "Which airspace classes need prior ATC authorization?",
    "How often must you complete recurrent training to keep flying under Part 107?",
    "Can a waiver allow flight over people? Which section says so?",
    "When is a drone under 0.55 lb exempt from registration?",
    "On a BYU home football game day, how far from the stadium and for how long is drone flight prohibited?",
]

# Part 4 - METAR groups to decode, in the order they appear
METAR_GROUPS = [
    "Report type and station (METAR or SPECI, KPVU)",
    "Day and time (Z)  -  convert to local time",
    "Wind (direction, speed, gusts)",
    "Visibility (SM)",
    "Present weather (RA, BR, FG, ...), if any",
    "Clouds (FEW / SCT / BKN / OVC and height in hundreds of ft AGL, or CLR)",
    "Temperature / dew point (C)",
    "Altimeter (A....)",
    "Remarks (RMK), if any",
]

GO_NOGO = [
    ("Flight visibility at least 3 SM?  (107.51(c))", ""),
    ("Ceiling: lowest BKN or OVC layer, in ft AGL  (CLR or no BKN/OVC = none)", ""),
    ("Could you stay 500 ft below the cloud at your planned altitude?  (107.51(d))", ""),
    ("Wind within the aircraft's limit?  Write the limit and its source", ""),
    ("TFRs over Provo at tfr.faa.gov?  Write 'none' or the NOTAM", ""),
    ("Airspace: KPVU is Class D. What authorization would a real flight need?", ""),
    ("GO or NO-GO for a Part 107 flight right now?", ""),
    ("One-sentence reason", ""),
]

STUDY_DAYS = 14


def title(ws, text, note=None, width_cols=5):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=BLUE)
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(italic=True, size=9, color="595959")
        ws["A2"].alignment = WRAP
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
        ws.row_dimensions[2].height = 30


def header_row(ws, row, names):
    for col, name in enumerate(names, start=1):
        c = ws.cell(row, col, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = 32


def label_cell(ws, row, col, text, fill=LABEL_FILL, bold=True):
    c = ws.cell(row, col, text)
    c.font = Font(bold=bold)
    c.fill = fill
    c.alignment = WRAP
    c.border = BOX
    return c


def blank_cells(ws, row, cols):
    for col in cols:
        c = ws.cell(row, col)
        c.alignment = WRAP
        c.border = BOX


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_part1(wb):
    ws = wb.active
    ws.title = "1 Scavenger Hunt"
    title(ws, "Part 107 Studying Lab - Resource Passport",
          "Fill in one copy per pair. Every answer needs the page or URL where you found it.")
    for r, label in enumerate(("Name", "Partner", "Date"), start=4):
        label_cell(ws, r, 1, label)
        c = ws.cell(r, 2)
        c.border = BOX
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws["A8"] = "Part 1 - Resource scavenger hunt (25 min)"
    ws["A8"].font = Font(bold=True, size=12, color=BLUE)
    hdr = 9
    header_row(ws, hdr, ["#", "Resource (open it)", "What to write", "Your answer", "Page or URL where you found it"])
    for i, (resource, what) in enumerate(FINDS, start=1):
        r = hdr + i
        label_cell(ws, r, 1, i)
        label_cell(ws, r, 2, resource, bold=False)
        label_cell(ws, r, 3, what, bold=False)
        blank_cells(ws, r, (4, 5))
        ws.row_dimensions[r].height = 42
    widths(ws, [5, 38, 40, 34, 40])
    ws.freeze_panes = ws.cell(hdr + 1, 1)


def sheet_part2(wb):
    ws = wb.create_sheet("2 Figure Questions")
    title(ws, "Part 2 - Reading a figure question (25 min)",
          "The question text is on the lab page. Write the figure and area you used, your answer, and why. "
          "Check against the answer page after the lab, not before.")
    hdr = 4
    header_row(ws, hdr, ["Question", "FAA sample question", "Figure and area used", "Your answer (A, B or C)", "Why (one sentence)"])
    for i, (label, q, fig) in enumerate(FIGURE_QS, start=1):
        r = hdr + i
        label_cell(ws, r, 1, label)
        label_cell(ws, r, 2, q, bold=False)
        label_cell(ws, r, 3, fig, bold=False)
        blank_cells(ws, r, (4, 5))
        ws.row_dimensions[r].height = 60
    widths(ws, [12, 22, 24, 22, 50])
    ws.freeze_panes = ws.cell(hdr + 1, 1)


def sheet_part3(wb):
    ws = wb.create_sheet("3 Look-It-Up Race")
    title(ws, "Part 3 - Look-it-up race (20 min)",
          "A right answer without a source scores half. Source = CFR section (for example 14 CFR 107.51(c)), "
          "AC 107-2A chapter and page, or the URL.")
    hdr = 4
    header_row(ws, hdr, ["#", "Question", "Your answer", "Source (section, chapter and page, or URL)"])
    for i, q in enumerate(RACE, start=1):
        r = hdr + i
        label_cell(ws, r, 1, i)
        label_cell(ws, r, 2, q, bold=False)
        blank_cells(ws, r, (3, 4))
        ws.row_dimensions[r].height = 48
    widths(ws, [5, 56, 36, 40])
    ws.freeze_panes = ws.cell(hdr + 1, 1)


def sheet_part4(wb):
    ws = wb.create_sheet("4 Live Weather")
    title(ws, "Part 4 - Today's real weather at KPVU (15 min)",
          "Paste the raw METAR from aviationweather.gov/data/metar/?ids=KPVU, then decode it one group at a time.",
          width_cols=3)
    label_cell(ws, 4, 1, "Time you pulled it (local)")
    blank_cells(ws, 4, (2, 3))
    label_cell(ws, 5, 1, "Raw METAR (paste the whole line)")
    blank_cells(ws, 5, (2, 3))
    ws.merge_cells("B5:C5")
    ws.row_dimensions[5].height = 45

    hdr = 7
    header_row(ws, hdr, ["Group", "As written in the METAR", "What it means"])
    for i, g in enumerate(METAR_GROUPS, start=1):
        r = hdr + i
        label_cell(ws, r, 1, g, bold=False)
        blank_cells(ws, r, (2, 3))
        ws.row_dimensions[r].height = 36

    r = hdr + len(METAR_GROUPS) + 2
    ws.cell(r, 1, "Could we fly under Part 107 right now?").font = Font(bold=True, size=12, color=BLUE)
    header_row(ws, r + 1, ["Check", "Your finding", "Notes"])
    for i, (check, _) in enumerate(GO_NOGO, start=1):
        rr = r + 1 + i
        label_cell(ws, rr, 1, check, bold=False)
        blank_cells(ws, rr, (2, 3))
        ws.row_dimensions[rr].height = 36
    widths(ws, [52, 34, 44])


def sheet_part5(wb):
    ws = wb.create_sheet("5 My Study Plan")
    title(ws, "Part 5 - My two-week study plan (15 min)",
          "Pick your two weakest test areas from the Drill-by-topic table on the Study Resources page. "
          "Name the exact resource for every day (a chapter, a figure, a video), not 'study charts'.",
          width_cols=6)
    for r, label in enumerate(("Weakest area 1 (and the resource for it)", "Weakest area 2 (and the resource for it)",
                               "Test date (or target date)"), start=4):
        label_cell(ws, r, 1, label)
        blank_cells(ws, r, (2,))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 30
    hdr = 8
    header_row(ws, hdr, ["Day", "Date", "Test area", "Resource (name it exactly)", "Minutes", "Done?"])
    for d in range(1, STUDY_DAYS + 1):
        r = hdr + d
        label_cell(ws, r, 1, d)
        blank_cells(ws, r, range(2, 7))
        ws.row_dimensions[r].height = 30
    widths(ws, [6, 14, 24, 52, 10, 8])
    ws.freeze_panes = ws.cell(hdr + 1, 1)


def build():
    wb = Workbook()
    sheet_part1(wb)
    sheet_part2(wb)
    sheet_part3(wb)
    sheet_part4(wb)
    sheet_part5(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
