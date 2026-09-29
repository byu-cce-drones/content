"""Build the downloadable Mission Log spreadsheet for the Creating Flight Plans lab.

Writes docs/labs/files/lab03_mission_log.xlsx. The rows must match the Mission Log table on
docs/labs/3_creating_flight_plans.md; edit both together.

    python fig_tools/make_mission_log.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "docs" / "labs" / "files" / "lab03_mission_log.xlsx"

MISSIONS = [
    "Mission 1 – Stadium",
    "Mission 2 – The Y",
    "Mission 3 – Utah Lake",
    "Mission 4 – Rock Canyon (Team)",
]

# (label, tall) - tall rows are for sentences rather than numbers.
ROWS = [
    ("Mission name (file name)", False),
    ("Drone and planning tool", False),
    ("Hard constraints", True),
    ("Mission objective", True),
    ("Relative altitude (ft)", False),
    ("Est. ground GSD (cm/px)", False),
    ("Frontal overlap (%)", False),
    ("Side overlap (%)", False),
    ("Flight speed (mph)", False),
    ("Gimbal pitch (°)", False),
    ("Passes", False),
    ("Total path distance (ft)", False),
    ("Estimated photos", False),
    ("Flight time (min, your calculation)", False),
    ("Batteries needed", False),
    ("All hard constraints met? (Y/N)", False),
    ("Most important decision, and why", True),
]

# Spreadsheet check of the flight time, from path distance and speed (1 mph = 88 ft/min).
CHECK_LABEL = "Flight time check (min) – calculated by the sheet"

BLUE = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=BLUE)
LABEL_FILL = PatternFill("solid", fgColor="DDEBF7")
CHECK_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Mission Log"

    ws["A1"] = "Creating Flight Plans Lab – Mission Log"
    ws["A1"].font = Font(bold=True, size=14, color=BLUE)
    ws["A3"] = ("Units follow the class flight planner: feet, miles per hour, and GSD in cm/px. "
                "If a controller app uses metres, note it in the cell.")
    ws["A3"].font = Font(italic=True, size=9, color="595959")
    ws.merge_cells("A3:E3")

    header_row = 5
    ws.cell(header_row, 1, "")
    for col, name in enumerate(MISSIONS, start=2):
        c = ws.cell(header_row, col, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BOX
    ws.cell(header_row, 1).fill = HEADER_FILL
    ws.cell(header_row, 1).border = BOX
    ws.row_dimensions[header_row].height = 32

    row = header_row + 1
    path_row = speed_row = None
    for label, tall in ROWS:
        c = ws.cell(row, 1, label)
        c.font = Font(bold=True)
        c.fill = LABEL_FILL
        c.alignment = WRAP
        c.border = BOX
        for col in range(2, 2 + len(MISSIONS)):
            cell = ws.cell(row, col)
            cell.alignment = WRAP
            cell.border = BOX
        ws.row_dimensions[row].height = 60 if tall else 20
        if label.startswith("Total path distance"):
            path_row = row
        if label.startswith("Flight speed"):
            speed_row = row
        row += 1

    c = ws.cell(row, 1, CHECK_LABEL)
    c.font = Font(italic=True)
    c.fill = CHECK_FILL
    c.alignment = WRAP
    c.border = BOX
    for col in range(2, 2 + len(MISSIONS)):
        L = get_column_letter(col)
        p, s = f"{L}{path_row}", f"{L}{speed_row}"
        cell = ws.cell(row, col, f'=IF(AND(ISNUMBER({p}),ISNUMBER({s}),{s}>0),ROUND({p}/({s}*88),1),"")')
        cell.fill = CHECK_FILL
        cell.border = BOX
        cell.alignment = WRAP
    ws.row_dimensions[row].height = 30

    ws.column_dimensions["A"].width = 34
    for col in range(2, 2 + len(MISSIONS)):
        ws.column_dimensions[get_column_letter(col)].width = 28
    ws.freeze_panes = ws.cell(header_row + 1, 2)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
